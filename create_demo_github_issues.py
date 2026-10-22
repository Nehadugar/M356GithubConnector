"""Create GitHub issues in GITHUB_REPO from the demo Excel workbook.

Commands:
    python create_demo_github_issues.py --dry-run          # preview every row; no GitHub writes
    python create_demo_github_issues.py                    # create the issues that do not exist yet
    python create_demo_github_issues.py --file other.xlsx  # use a different workbook

Rows whose title already exists in the repository (open or closed) are skipped, so the script
is safe to re-run. Rows with Status "closed" are created and then closed.
"""

import argparse
import sys
import time
from dataclasses import dataclass
from pathlib import Path

import requests
from dotenv import load_dotenv
from openpyxl import load_workbook

from github_issues_connector import GITHUB_API, ConnectorError, env, github_repo, require_env

DEFAULT_FILE = Path(__file__).with_name("Copilot_GitHub_Issues_Demo_Dataset.xlsx")
SHEET = "Issues"
HEADER_MARKER = "Issue #"
REQUIRED_COLUMNS = ("Issue #", "Title", "Description", "Status", "Labels")
MAX_TITLE_LEN = 256
MAX_LABEL_LEN = 50
MAX_ATTEMPTS = 4
CREATE_INTERVAL_S = 1.0  # GitHub asks for no more than about one content-creating request per second

ACTION_CREATE = "create"
ACTION_CREATE_AND_CLOSE = "create and close"
ACTION_SKIP_EXISTING = "skip (exists)"
ACTION_SKIP_DUPLICATE = "skip (duplicate in workbook)"


@dataclass
class DemoIssue:
    row: int
    demo_number: str
    title: str
    description: str
    status: str  # "open" or "closed"
    labels: list[str]
    author: str


@dataclass
class PlanItem:
    issue: DemoIssue
    action: str
    existing: dict | None = None  # the matching GitHub issue for skip (exists)


def normalize_title(title: str) -> str:
    return " ".join(title.split()).casefold()


def cell_text(ws, row: int, column: int) -> str:
    value = ws.cell(row, column).value
    return "" if value is None else str(value).strip()


def read_workbook(path: Path) -> list[DemoIssue]:
    """Read and validate every issue row. Raises before any GitHub write if a row is invalid."""
    try:
        workbook = load_workbook(path, data_only=True)
    except FileNotFoundError:
        raise ConnectorError(f"Workbook not found: {path}")
    if SHEET not in workbook.sheetnames:
        raise ConnectorError(f"Sheet '{SHEET}' not found in {path.name}")
    ws = workbook[SHEET]

    header_row, columns = None, {}
    for r in range(1, 21):
        found = {cell_text(ws, r, c): c for c in range(1, ws.max_column + 1) if cell_text(ws, r, c)}
        if HEADER_MARKER in found:
            header_row, columns = r, found
            break
    if header_row is None:
        raise ConnectorError(f"Could not find the header row (column '{HEADER_MARKER}') in sheet '{SHEET}'.")
    missing = [name for name in REQUIRED_COLUMNS if name not in columns]
    if missing:
        raise ConnectorError(f"Missing column(s) in sheet '{SHEET}': {', '.join(missing)}")

    issues, errors = [], []
    for r in range(header_row + 1, ws.max_row + 1):
        title = cell_text(ws, r, columns["Title"])
        if not title:
            continue
        status = cell_text(ws, r, columns["Status"]).lower()
        if status not in ("open", "closed"):
            errors.append(f"row {r}: Status must be 'open' or 'closed', got '{status}'")
        if len(title) > MAX_TITLE_LEN:
            errors.append(f"row {r}: title is longer than {MAX_TITLE_LEN} characters")
        labels = [label.strip() for label in cell_text(ws, r, columns["Labels"]).split(",") if label.strip()]
        for label in labels:
            if len(label) > MAX_LABEL_LEN:
                errors.append(f"row {r}: label '{label}' is longer than {MAX_LABEL_LEN} characters")
        issues.append(
            DemoIssue(
                row=r,
                demo_number=cell_text(ws, r, columns["Issue #"]),
                title=title,
                description=cell_text(ws, r, columns["Description"]),
                status=status,
                labels=labels,
                author=cell_text(ws, r, columns["Author"]) if "Author" in columns else "",
            )
        )

    if errors:
        raise ConnectorError("Workbook validation failed:\n  " + "\n  ".join(errors))
    if not issues:
        raise ConnectorError(f"No issues found in sheet '{SHEET}'.")
    return issues


class GitHubAccessError(ConnectorError):
    """The token cannot use the repository, so later rows would fail the same way."""


ACCESS_HINT = (
    "\nHint: GITHUB_TOKEN needs write access to Issues on this repository. For a fine-grained token, "
    "edit it in GitHub settings and set Repository permissions > Issues to 'Read and write', "
    "then update GITHUB_TOKEN in .env."
)


def raise_github_error(action: str, resp: requests.Response):
    try:
        message = resp.json().get("message", resp.text)
    except ValueError:
        message = resp.text
    text = f"{action} failed: HTTP {resp.status_code} - {message}"
    if resp.status_code in (401, 403) and "rate limit" not in resp.text.lower():
        raise GitHubAccessError(text + ACCESS_HINT)
    raise ConnectorError(text)


def github_call(session: requests.Session, method: str, url: str, **kwargs) -> requests.Response:
    """Send a request, retrying on rate limits and transient server errors."""
    for attempt in range(1, MAX_ATTEMPTS + 1):
        resp = session.request(method, url, timeout=60, **kwargs)
        rate_limited = resp.status_code == 403 and "rate limit" in resp.text.lower()
        retryable = resp.status_code in (429, 502, 503, 504) or rate_limited
        if not retryable or attempt == MAX_ATTEMPTS:
            return resp
        time.sleep(int(resp.headers.get("Retry-After", 30 * attempt)))
    return resp


def fetch_existing_issues(session: requests.Session, repo: str) -> dict[str, dict]:
    """Map normalized title -> GitHub issue for every open and closed issue (pull requests excluded)."""
    url = f"{GITHUB_API}/repos/{repo}/issues"
    params = {"state": "all", "per_page": 100}
    existing: dict[str, dict] = {}
    while url:
        resp = github_call(session, "GET", url, params=params)
        params = None  # the "next" link already carries the query string
        if not resp.ok:
            raise_github_error("List existing issues", resp)
        for issue in resp.json():
            if "pull_request" in issue:
                continue
            existing.setdefault(normalize_title(issue["title"]), issue)
        url = resp.links.get("next", {}).get("url")
    return existing


def build_plan(issues: list[DemoIssue], existing: dict[str, dict]) -> list[PlanItem]:
    plan, seen = [], set()
    for issue in issues:
        key = normalize_title(issue.title)
        if key in seen:
            plan.append(PlanItem(issue, ACTION_SKIP_DUPLICATE))
            continue
        seen.add(key)
        if key in existing:
            plan.append(PlanItem(issue, ACTION_SKIP_EXISTING, existing[key]))
        elif issue.status == "closed":
            plan.append(PlanItem(issue, ACTION_CREATE_AND_CLOSE))
        else:
            plan.append(PlanItem(issue, ACTION_CREATE))
    return plan


def build_body(issue: DemoIssue) -> str:
    footer = f"---\nSample issue from the demo workbook (demo ref #{issue.demo_number}). Original author: {issue.author or 'n/a'}."
    return f"{issue.description}\n\n{footer}" if issue.description else footer


def create_issue(session: requests.Session, repo: str, issue: DemoIssue) -> dict:
    payload = {"title": issue.title, "body": build_body(issue), "labels": issue.labels}
    resp = github_call(session, "POST", f"{GITHUB_API}/repos/{repo}/issues", json=payload)
    if not resp.ok:
        raise_github_error("Create issue", resp)
    created = resp.json()
    if issue.status == "closed":
        url = f"{GITHUB_API}/repos/{repo}/issues/{created['number']}"
        resp = github_call(session, "PATCH", url, json={"state": "closed", "state_reason": "completed"})
        if not resp.ok:
            raise ConnectorError(
                f"Close issue failed: HTTP {resp.status_code} - {resp.text}\n"
                f"Issue #{created['number']} was created but is still open: {created['html_url']}"
            )
    return created


def print_plan(plan: list[PlanItem], repo: str):
    print(f"Target repository: {repo}")
    print("Dry run: no GitHub changes will be made.\n")
    print(f"{'Row':<5}{'Demo #':<8}{'Status':<8}{'Action':<30}{'Title':<60}Labels")
    for item in plan:
        i = item.issue
        action = item.action
        if item.existing:
            action = f"{action} #{item.existing['number']}"
        title = i.title if len(i.title) <= 58 else i.title[:55] + "..."
        print(f"{i.row:<5}{i.demo_number:<8}{i.status:<8}{action:<30}{title:<60}{', '.join(i.labels)}")
    counts = {}
    for item in plan:
        counts[item.action] = counts.get(item.action, 0) + 1
    print("\nSummary: " + ", ".join(f"{count} x {action}" for action, count in counts.items()))


def run(args) -> int:
    token = env("GITHUB_TOKEN")
    require_env("GITHUB_TOKEN")
    repo = github_repo()
    issues = read_workbook(args.file)

    session = requests.Session()
    session.headers.update(
        {
            "Accept": "application/vnd.github+json",
            "X-GitHub-Api-Version": "2022-11-28",
            "Authorization": f"Bearer {token}",
        }
    )
    existing = fetch_existing_issues(session, repo)
    plan = build_plan(issues, existing)

    if args.dry_run:
        print_plan(plan, repo)
        return 0

    created, skipped, failed = [], [], []
    print(f"Creating issues in {repo} from {args.file.name}...")
    for item in plan:
        i = item.issue
        if item.action in (ACTION_SKIP_EXISTING, ACTION_SKIP_DUPLICATE):
            skipped.append(item)
            continue
        try:
            result = create_issue(session, repo, i)
            created.append((i, result))
            print(f"  [OK] row {i.row}: #{result['number']} ({i.status}) {result['html_url']}")
        except GitHubAccessError:
            raise
        except (ConnectorError, requests.RequestException) as e:
            failed.append((i, str(e)))
            print(f"  [FAILED] row {i.row}: {i.title}\n    {e}", file=sys.stderr)
        time.sleep(CREATE_INTERVAL_S)

    print("\nCreated issues:")
    for i, result in created:
        print(f"  #{result['number']} [{i.status}] {i.title}\n    {result['html_url']}")
    if skipped:
        print("\nSkipped (already exist or duplicate in workbook):")
        for item in skipped:
            where = f" as #{item.existing['number']} ({item.existing['html_url']})" if item.existing else ""
            print(f"  row {item.issue.row}: {item.issue.title}{where}")
    if failed:
        print("\nFailed:")
        for i, message in failed:
            print(f"  row {i.row}: {i.title}\n    {message}")
    print(f"\nDone: {len(created)} created, {len(skipped)} skipped, {len(failed)} failed.")
    return 1 if failed else 0


def main() -> int:
    sys.stdout.reconfigure(errors="replace")
    parser = argparse.ArgumentParser(description="Create GitHub issues from the demo Excel workbook.")
    parser.add_argument("--file", type=Path, default=DEFAULT_FILE, help="workbook path (default: the demo dataset)")
    parser.add_argument("--dry-run", action="store_true", help="preview all issues without creating anything")
    args = parser.parse_args()

    load_dotenv()
    try:
        return run(args)
    except ConnectorError as e:
        print(f"Error: {e}", file=sys.stderr)
        return 1
    except requests.RequestException as e:
        print(f"Network error: {e}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
